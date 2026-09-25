from __future__ import annotations

import asyncio
import csv
import io
import ipaddress
import mimetypes
import os
import socket
import uuid
from pathlib import Path
from urllib.parse import urlparse

import httpx
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
AVATAR_DIR = BASE_DIR / "avatar_cache"
AVATAR_DIR.mkdir(exist_ok=True)
MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_AVATAR_SIZE = 4 * 1024 * 1024
MAX_ROWS = 10_000
AVATAR_COLORS = ["#e7c8b3", "#c7d5c3", "#e7c2c7", "#c9d2e0", "#dfd2a7", "#c4d8dc"]

app = FastAPI(title="Lucky Draw Import API")
cors_origins = [
    origin.strip()
    for origin in os.getenv(
        "CORS_ORIGINS",
        "http://localhost:4200,http://127.0.0.1:4200",
    ).split(",")
    if origin.strip()
]
app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_origins,
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
)
app.mount("/avatars", StaticFiles(directory=AVATAR_DIR), name="avatars")


def normalize_header(value: object) -> str:
    return "".join(character for character in str(value or "").lower() if character.isalnum())


def parse_csv(content: bytes) -> list[list[object]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise HTTPException(status_code=400, detail="The CSV must be encoded as UTF-8.") from error

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def parse_xlsx(content: bytes) -> list[list[object]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        workbook.close()
        return rows
    except Exception as error:
        raise HTTPException(status_code=400, detail="The Excel file could not be read. Use a valid .xlsx workbook.") from error


def get_accounts(rows: list[list[object]]) -> list[dict[str, str]]:
    if not rows:
        raise HTTPException(status_code=400, detail="The file is empty.")

    headers = [normalize_header(cell) for cell in rows[0]]
    username_headers = {"username", "instagramusername", "instagram", "handle", "account"}
    avatar_headers = {"avatarurl", "profilepictureurl", "profilepicture", "profilepic", "imageurl", "avatar"}
    username_index = next((index for index, name in enumerate(headers) if name in username_headers), None)
    avatar_index = next((index for index, name in enumerate(headers) if name in avatar_headers), None)

    if username_index is None:
        raise HTTPException(
            status_code=400,
            detail="A username column is required. Use username, instagram_username, instagram, or handle.",
        )

    accounts: list[dict[str, str]] = []
    seen_usernames: set[str] = set()
    for row in rows[1:]:
        username = str(row[username_index] or "").strip().lstrip("@").strip()
        if not username:
            continue
        username_key = username.casefold()
        if username_key in seen_usernames:
            continue
        seen_usernames.add(username_key)

        avatar_url = str(row[avatar_index] or "").strip() if avatar_index is not None and avatar_index < len(row) else ""
        accounts.append({"username": username, "avatar_url": avatar_url})
        if len(accounts) > MAX_ROWS:
            raise HTTPException(status_code=400, detail=f"A file can contain at most {MAX_ROWS} accounts.")

    if not accounts:
        raise HTTPException(status_code=400, detail="No usernames were found below the header row.")
    return accounts


def is_public_host(hostname: str) -> bool:
    try:
        addresses = {ipaddress.ip_address(hostname)}
    except ValueError:
        try:
            addresses = {
                ipaddress.ip_address(result[4][0])
                for result in socket.getaddrinfo(hostname, None, type=socket.SOCK_STREAM)
            }
        except OSError:
            return False
    return bool(addresses) and all(address.is_global for address in addresses)


async def download_avatar(client: httpx.AsyncClient, url: str, index: int) -> tuple[str, str | None]:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or not is_public_host(parsed.hostname):
        return "", f"Row {index}: avatar URL must point to a public HTTP or HTTPS image."

    try:
        async with client.stream("GET", url) as response:
            response.raise_for_status()
            content_type = response.headers.get("content-type", "").split(";", 1)[0].lower()
            extension = mimetypes.guess_extension(content_type)
            if not content_type.startswith("image/") or extension not in {".jpg", ".jpeg", ".png", ".webp", ".gif"}:
                return "", f"Row {index}: avatar URL did not return a supported image."

            image = bytearray()
            async for chunk in response.aiter_bytes():
                image.extend(chunk)
                if len(image) > MAX_AVATAR_SIZE:
                    return "", f"Row {index}: avatar image exceeded the 4 MB limit."

        filename = f"{uuid.uuid4().hex}{extension}"
        (AVATAR_DIR / filename).write_bytes(image)
        return f"/avatars/{filename}", None
    except (httpx.HTTPError, OSError, ValueError):
        return "", f"Row {index}: avatar could not be downloaded."


@app.get("/api/health")
async def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/api/import")
async def import_entries(file: UploadFile = File(...)) -> dict[str, object]:
    filename = file.filename or ""
    extension = Path(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise HTTPException(status_code=400, detail="Choose a .csv or .xlsx file.")

    content = await file.read(MAX_FILE_SIZE + 1)
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="The file must be smaller than 10 MB.")

    rows = parse_csv(content) if extension == ".csv" else parse_xlsx(content)
    accounts = get_accounts(rows)
    timeout = httpx.Timeout(8.0, connect=4.0)
    limits = httpx.Limits(max_connections=12, max_keepalive_connections=6)
    semaphore = asyncio.Semaphore(6)

    async with httpx.AsyncClient(timeout=timeout, limits=limits, follow_redirects=False) as client:
        async def process_account(index: int, account: dict[str, str]) -> tuple[dict[str, object], str | None]:
            avatar = ""
            warning = None
            if account["avatar_url"]:
                async with semaphore:
                    avatar, warning = await download_avatar(client, account["avatar_url"], index + 2)
            entrant = {
                "username": account["username"],
                "avatar": avatar,
                "color": AVATAR_COLORS[index % len(AVATAR_COLORS)],
            }
            return entrant, warning

        processed = await asyncio.gather(*(process_account(index, account) for index, account in enumerate(accounts)))

    entrants = [entrant for entrant, _ in processed]
    warnings = [warning for _, warning in processed if warning]
    if len(warnings) > 5:
        warnings = warnings[:5] + [f"and {len(warnings) - 5} more avatar issue(s)."]
    return {"entrants": entrants, "warnings": warnings}
